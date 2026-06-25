import flet as ft
import sqlite3
import os
import random
from datetime import date, timedelta

APP_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(APP_DIR, "freddy_fazbear_mobile.db")

# ──────────────────────── База данных ────────────────────────
class Database:
    def __init__(self):
        self.conn = sqlite3.connect(DB_PATH)
        self.conn.row_factory = sqlite3.Row
        self.create_table()

    def create_table(self):
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS tickets (
                id TEXT PRIMARY KEY,
                attraction TEXT NOT NULL,
                visitor_name TEXT NOT NULL,
                age INTEGER NOT NULL,
                price REAL NOT NULL,
                ticket_type TEXT NOT NULL,
                purchase_date TEXT NOT NULL
            )
        """)
        self.conn.commit()
        if self.count() == 0:
            self.seed()

    def count(self):
        return self.conn.execute("SELECT COUNT(*) FROM tickets").fetchone()[0]

    def seed(self):
        attractions = [
            "Караоке с Фредди", "Холм холостяков", "Аттракцион «Звёзды»",
            "Комната Бонни", "Кукольный театр Чика", "Ресторан «Пицца Фрэдди»",
            "Спиральная горка Фокси", "Комната страха", "Водопад Эндорфина",
            "Световое шоу", "Космический полёт", "Карусель мечты"
        ]
        names = [
            "Иван Петров", "Мария Сидорова", "Алексей Козлов", "Анна Морозова",
            "Дмитрий Новиков", "Елена Соколова", "Сергей Лебедев", "Ольга Попова",
            "Никита Волков", "Татьяна Зайцева", "Максим Фёдоров", "Наталья Михайлова"
        ]
        for i in range(12):
            att = random.choice(attractions)
            name = random.choice(names)
            age = random.randint(3, 60)
            ticket_type = "Детский" if age < 14 else "Взрослый"
            price = round(random.uniform(150, 999), 2)
            rid = f"АТ-{i:03d}"
            d = date(2025, 1, 1) + timedelta(days=random.randint(0, 364))
            self.conn.execute(
                "INSERT INTO tickets VALUES (?,?,?,?,?,?,?)",
                (rid, att, name, age, price, ticket_type, d.strftime("%d.%m.%Y"))
            )
        self.conn.commit()

    def get_next_id(self):
        rows = self.conn.execute("SELECT id FROM tickets").fetchall()
        nums = []
        for r in rows:
            try:
                nums.append(int(r["id"].split("-")[1]))
            except (IndexError, ValueError):
                pass
        nxt = max(nums) + 1 if nums else 0
        return f"АТ-{nxt:03d}"

    def fetch_all(self):
        return self.conn.execute("SELECT * FROM tickets ORDER BY id").fetchall()

    def add(self, attraction, visitor, age, price, ttype, purchase_date):
        self.conn.execute(
            "INSERT INTO tickets VALUES (?,?,?,?,?,?,?)",
            (self.get_next_id(), attraction, visitor, age, price, ttype, purchase_date)
        )
        self.conn.commit()

    def update(self, tid, attraction, visitor, age, price, ttype, purchase_date):
        self.conn.execute(
            "UPDATE tickets SET attraction=?, visitor_name=?, age=?, price=?, ticket_type=?, purchase_date=? WHERE id=?",
            (attraction, visitor, age, price, ttype, purchase_date, tid)
        )
        self.conn.commit()

    def delete(self, tid):
        self.conn.execute("DELETE FROM tickets WHERE id=?", (tid,))
        self.conn.commit()

    def query(self, sql, params=()):
        return self.conn.execute(sql, params).fetchall()


# ────────────────── Цвета ──────────────────
class Colors:
    BG = "#1a1a2e"
    BG_MID = "#16213e"
    BG_LIGHT = "#0f3460"
    FG = "#e0e0e0"
    FG_DIM = "#a0a0b0"
    ACCENT = "#e94560"
    ACCENT_HOVER = "#ff6b81"
    TREE_BG = "#0d1b2a"
    ENTRY_BG = "#1a1a3e"
    ERROR = "#ff6b6b"
    WHITE = "#ffffff"


COLUMNS = [
    ("id", "ID"),
    ("attraction", "Аттракцион"),
    ("visitor_name", "Посетитель"),
    ("age", "Возраст"),
    ("price", "Цена"),
    ("ticket_type", "Тип билета"),
    ("purchase_date", "Дата покупки"),
]

db = Database()


# ────────────────── Основное приложение ──────────────────
def main(page: ft.Page):
    page.title = "Freddy Fazbear"
    def show_snack(msg, bg="#2ecc71"):
        sb = ft.SnackBar(content=ft.Text(msg, color=ft.Colors.WHITE), bgcolor=bg, open=True)
        page.overlay.append(sb)
        page.update()
    page.bgcolor = Colors.BG
    page.theme_mode = ft.ThemeMode.DARK
    page.window.width = 420
    page.window.height = 800
    page.padding = 0
    page.spacing = 0

    visible_cols = [c[0] for c in COLUMNS]
    search_query = ""
    sort_col = None
    sort_rev = False

    # ─── Заголовок ───
    header = ft.Container(
        content=ft.Row(
            [
                ft.Text("FREDDY FAZBEAR", size=20, weight=ft.FontWeight.BOLD,
                        color=Colors.ACCENT),
                ft.Text("СУБД", size=14, color=Colors.FG_DIM),
            ],
            alignment=ft.MainAxisAlignment.START,
        ),
        bgcolor=Colors.BG_MID,
        padding=ft.padding.Padding(left=16, top=12, right=16, bottom=12),
        border_radius=ft.BorderRadius.only(bottom_left=12, bottom_right=12),
    )

    # ─── Статус-бар ───
    status_text = ft.Text("Готово", size=12, color=Colors.FG_DIM)
    count_text = ft.Text("", size=12, color=Colors.FG_DIM)

    # ─── Таблица ───
    data_table = ft.DataTable(
        columns=[],
        rows=[],
        heading_row_color=Colors.BG_LIGHT,
        heading_row_height=40,
        data_row_color={"": Colors.TREE_BG},
        data_row_min_height=38,
        data_row_max_height=38,
        horizontal_margin=8,
        column_spacing=8,
        divider_thickness=0.5,
        border=ft.Border(
            left=ft.BorderSide(0.5, Colors.BG_LIGHT),
            top=ft.BorderSide(0.5, Colors.BG_LIGHT),
            right=ft.BorderSide(0.5, Colors.BG_LIGHT),
            bottom=ft.BorderSide(0.5, Colors.BG_LIGHT),
        ),
        border_radius=8,
    )

    table_scroll = ft.Container(
        content=ft.Column(
            [data_table],
            scroll=ft.ScrollMode.AUTO,
        ),
        expand=True,
        padding=ft.padding.Padding(left=8, top=0, right=8, bottom=0),
    )

    def rebuild_table():
        data_table.columns = []
        for cid, heading in COLUMNS:
            if cid in visible_cols:
                w = 70 if cid == "id" else 120 if cid in ("attraction", "visitor_name") else 100
                data_table.columns.append(
                    ft.DataColumn(
                        label=ft.Text(heading, size=11, weight=ft.FontWeight.BOLD, color=Colors.FG),
                        on_sort=lambda e, c=cid: on_sort(c),
                    )
                )
        refresh()

    selected_row_id = None

    def refresh():
        data = db.fetch_all()
        if sort_col:
            idx = next(i for i, c in enumerate(COLUMNS) if c[0] == sort_col)
            data = sorted(data, key=lambda r: list(r)[idx], reverse=sort_rev)

        if search_query:
            q = search_query.lower()
            data = [r for r in data if any(q in str(list(r)[i]).lower() for i in range(len(COLUMNS)))]

        visible_indices = [i for i, c in enumerate(COLUMNS) if c[0] in visible_cols]

        data_table.rows = []
        for row in data:
            cells = []
            for vi in visible_indices:
                val = str(list(row)[vi])
                cells.append(
                    ft.DataCell(
                        ft.Text(val, size=10, color=Colors.FG,
                                overflow=ft.TextOverflow.ELLIPSIS, max_lines=1)
                    )
                )
            row_id = list(row)[0]
            is_selected = row_id == selected_row_id
            data_table.rows.append(
                ft.DataRow(
                    cells=cells,
                    on_select_change=lambda e, r=row: on_row_select(list(r)),
                    color={"": Colors.ACCENT if is_selected else Colors.TREE_BG},
                    selected=is_selected,
                )
            )

        count_text.value = f"Записей: {len(data)}"
        status_text.value = "Данные обновлены"
        page.update()

    def on_sort(col):
        nonlocal sort_col, sort_rev
        if sort_col == col:
            sort_rev = not sort_rev
        else:
            sort_col = col
            sort_rev = False
        refresh()

    # ─── Поиск ───
    search_field = ft.TextField(
        hint_text="Поиск по всем столбцам...",
        prefix_icon=ft.Icons.SEARCH,
        bgcolor=Colors.ENTRY_BG,
        color=Colors.FG,
        hint_style=ft.TextStyle(color=Colors.FG_DIM),
        border_color=Colors.BG_LIGHT,
        focused_border_color=Colors.ACCENT,
        expand=True,
        height=44,
        text_size=13,
        on_submit=lambda e: do_search(),
        content_padding=ft.padding.Padding(left=8, top=12, right=8, bottom=12),
    )

    def do_search():
        nonlocal search_query
        search_query = search_field.value.strip() if search_field.value else ""
        refresh()

    def clear_search():
        nonlocal search_query
        search_field.value = ""
        search_query = ""
        refresh()

    search_bar = ft.Container(
        content=ft.Row(
            [
                search_field,
                ft.IconButton(ft.Icons.CLOSE, icon_color=Colors.FG_DIM,
                              on_click=lambda _: clear_search(), icon_size=20),
            ],
            vertical_alignment=ft.CrossAxisAlignment.CENTER,
        ),
        padding=ft.padding.Padding(left=8, top=8, right=8, bottom=8),
    )

    # ─── Фильтр столбцов ───
    col_filter_field = ft.TextField(
        hint_text="ID, Аттракцион, Посетитель...",
        bgcolor=Colors.ENTRY_BG,
        color=Colors.FG,
        hint_style=ft.TextStyle(color=Colors.FG_DIM),
        border_color=Colors.BG_LIGHT,
        focused_border_color=Colors.ACCENT,
        expand=True,
        height=44,
        text_size=13,
        on_submit=lambda e: apply_col_filter(),
        content_padding=ft.padding.Padding(left=8, top=12, right=8, bottom=12),
    )

    def apply_col_filter():
        raw = col_filter_field.value.strip() if col_filter_field.value else ""
        nonlocal visible_cols
        if not raw:
            visible_cols = [c[0] for c in COLUMNS]
        else:
            requested = [s.strip() for s in raw.split(",") if s.strip()]
            visible_cols = []
            for cid, heading in COLUMNS:
                if heading in requested or cid in requested:
                    visible_cols.append(cid)
            if not visible_cols:
                visible_cols = [c[0] for c in COLUMNS]
        rebuild_table()

    def reset_col_filter():
        nonlocal visible_cols
        col_filter_field.value = ""
        visible_cols = [c[0] for c in COLUMNS]
        rebuild_table()

    col_names = ", ".join(c[1] for c in COLUMNS)
    col_filter_bar = ft.Container(
        content=ft.Column(
            [
                ft.Text(f"Столбцы: {col_names}", size=10, color=Colors.FG_DIM),
                ft.Row(
                    [
                        col_filter_field,
                        ft.IconButton(ft.Icons.REFRESH, icon_color=Colors.ACCENT,
                                      on_click=lambda _: reset_col_filter(), icon_size=20),
                    ],
                ),
            ],
            spacing=2,
        ),
        padding=ft.padding.Padding(left=8, top=4, right=8, bottom=4),
    )

    # ─── Кнопки CRUD ───
    def open_add_dialog(_=None):
        open_dialog("Добавить запись", None)

    selected_record = None

    def on_row_select(row_data):
        nonlocal selected_record, selected_row_id
        selected_record = row_data
        selected_row_id = row_data[0]
        refresh()
        open_dialog("Изменить запись", row_data)

    def delete_selected(_=None):
        if not selected_record:
            show_snack("Выберите запись", Colors.ACCENT)
            return

        def confirm_delete(_):
            nonlocal selected_record, selected_row_id
            try:
                db.delete(selected_record[0])
                selected_record = None
                selected_row_id = None
                refresh()
                page.pop_dialog()
                show_snack(f"Запись удалена", Colors.ACCENT)
            except Exception as ex:
                page.pop_dialog()
                show_snack(f"Ошибка: {ex}", Colors.ERROR)

        dlg = ft.AlertDialog(
            title=ft.Text("Подтверждение"),
            content=ft.Text(f"Удалить запись {selected_record[0]}?"),
            actions=[
                ft.TextButton("Да", on_click=confirm_delete),
                ft.TextButton("Нет", on_click=lambda _: page.pop_dialog()),
            ],
        )
        page.show_dialog(dlg)

    # ─── Диалог добавления/изменения ───
    def open_dialog(title, record):
        id_val = record[0] if record else db.get_next_id()
        attraction_val = record[1] if record else ""
        visitor_val = record[2] if record else ""
        age_val = str(record[3]) if record else ""
        price_val = str(record[4]) if record else ""
        type_val = record[5] if record else "Взрослый"
        date_val = record[6] if record else date.today().strftime("%d.%m.%Y")

        tf_id = ft.TextField(value=id_val, label="ID", read_only=True, bgcolor=Colors.ENTRY_BG,
                             color=Colors.FG, border_color=Colors.BG_LIGHT, text_size=13, height=48)
        tf_att = ft.TextField(value=attraction_val, label="Аттракцион", bgcolor=Colors.ENTRY_BG,
                              color=Colors.FG, border_color=Colors.BG_LIGHT, text_size=13, height=48)
        tf_vis = ft.TextField(value=visitor_val, label="Посетитель", bgcolor=Colors.ENTRY_BG,
                              color=Colors.FG, border_color=Colors.BG_LIGHT, text_size=13, height=48)
        tf_age = ft.TextField(value=age_val, label="Возраст", bgcolor=Colors.ENTRY_BG,
                              color=Colors.FG, border_color=Colors.BG_LIGHT, text_size=13, height=48,
                              keyboard_type=ft.KeyboardType.NUMBER)
        tf_price = ft.TextField(value=price_val, label="Цена", bgcolor=Colors.ENTRY_BG,
                                color=Colors.FG, border_color=Colors.BG_LIGHT, text_size=13, height=48,
                                keyboard_type=ft.KeyboardType.NUMBER)
        dd_type = ft.Dropdown(
            value=type_val, label="Тип билета",
            options=[ft.dropdown.Option("Детский"), ft.dropdown.Option("Взрослый")],
            bgcolor=Colors.ENTRY_BG, color=Colors.FG,
            border_color=Colors.BG_LIGHT, focused_border_color=Colors.ACCENT,
            text_size=13, height=52,
        )
        tf_date = ft.TextField(value=date_val, label="Дата покупки (ДД.ММ.ГГГГ)", bgcolor=Colors.ENTRY_BG,
                               color=Colors.FG, border_color=Colors.BG_LIGHT, text_size=13, height=48)
        err_text = ft.Text("", color=Colors.ERROR, size=12)

        def auto_type(_):
            try:
                a = int(tf_age.value)
                dd_type.value = "Детский" if a <= 12 else "Взрослый"
                page.update()
            except ValueError:
                pass

        tf_age.on_change = auto_type
        if not record:
            auto_type(None)

        def save(_):
            attraction = tf_att.value.strip()
            visitor = tf_vis.value.strip()
            age_str = tf_age.value.strip()
            price_str = tf_price.value.strip()
            ttype = dd_type.value
            pdate = tf_date.value.strip()

            if not attraction or not visitor or not age_str or not price_str or not pdate:
                err_text.value = "Заполните все поля"
                page.update()
                return

            if any(ch.isdigit() for ch in attraction):
                err_text.value = "Аттракцион — только буквы"
                page.update()
                return
            if any(ch.isdigit() for ch in visitor):
                err_text.value = "Посетитель — только буквы"
                page.update()
                return

            try:
                age = int(age_str)
                if age < 0 or age > 120:
                    err_text.value = "Возраст 0–120"
                    page.update()
                    return
            except ValueError:
                err_text.value = "Возраст — целое число"
                page.update()
                return

            if ttype == "Детский" and age > 12:
                err_text.value = "Детский билет: до 12 лет"
                page.update()
                return
            if ttype == "Взрослый" and age < 13:
                err_text.value = "Взрослый билет: от 13 лет"
                page.update()
                return

            try:
                price = float(price_str)
                if price <= 0:
                    err_text.value = "Цена > 0"
                    page.update()
                    return
            except ValueError:
                err_text.value = "Цена — число"
                page.update()
                return

            try:
                parts = pdate.split(".")
                if len(parts) != 3:
                    raise ValueError
                d, m, y = int(parts[0]), int(parts[1]), int(parts[2])
                pd = date(y, m, d)
            except (ValueError, IndexError):
                err_text.value = "Дата — ДД.ММ.ГГГГ"
                page.update()
                return

            if pd > date.today():
                err_text.value = "Дата не позже сегодня"
                page.update()
                return

            try:
                if record:
                    db.update(id_val, attraction, visitor, age, price, ttype, pdate)
                else:
                    db.add(attraction, visitor, age, price, ttype, pdate)
                refresh()
                page.pop_dialog()
                show_snack(f"Запись {id_val} обновлена" if record else "Запись добавлена")
            except Exception as ex:
                page.pop_dialog()
                show_snack(f"Ошибка БД: {ex}", Colors.ERROR)

        dlg = ft.AlertDialog(
            title=ft.Text(title, size=18, weight=ft.FontWeight.BOLD, color=Colors.ACCENT),
            content=ft.Container(
                content=ft.Column(
                    [tf_id, tf_att, tf_vis, tf_age, tf_price, dd_type, tf_date, err_text],
                    spacing=4,
                    scroll=ft.ScrollMode.AUTO,
                ),
                width=340,
                height=420,
            ),
            actions=[ft.TextButton("Сохранить", on_click=save,
                                   style=ft.ButtonStyle(color=Colors.ACCENT)),
                     ft.TextButton("Отмена", on_click=lambda _: page.pop_dialog())],
            actions_alignment=ft.MainAxisAlignment.END,
        )
        page.show_dialog(dlg)

    # ─── Отчёты ───
    def open_reports(_=None):
        content_col = ft.Column(spacing=2, scroll=ft.ScrollMode.AUTO)
        expanded_cat = {"val": None}

        def rebuild():
            content_col.controls.clear()
            for cat_name, items in report_tree.items():
                is_open = expanded_cat["val"] == cat_name
                arrow = ft.Icon(ft.Icons.ARROW_DROP_DOWN if is_open else ft.Icons.ARROW_RIGHT,
                                color=Colors.ACCENT, size=18)

                def make_toggle(cn=cat_name):
                    def toggle(_):
                        expanded_cat["val"] = cn if expanded_cat["val"] != cn else None
                        rebuild()
                    return toggle

                content_col.controls.append(
                    ft.Container(
                        content=ft.Row([arrow, ft.Text(cat_name, size=12, weight=ft.FontWeight.BOLD,
                                                        color=Colors.ACCENT, expand=True)],
                                       vertical_alignment=ft.CrossAxisAlignment.CENTER),
                        bgcolor=Colors.BG_MID, border_radius=6,
                        padding=ft.padding.Padding(left=6, top=4, right=6, bottom=4),
                        on_click=make_toggle(), ink=True,
                    ))
                if is_open:
                    for label, func in items:
                        def make_click(f=func):
                            def click(_):
                                page.pop_dialog()
                                f()
                            return click
                        content_col.controls.append(
                            ft.Container(
                                content=ft.Row([
                                    ft.Icon(ft.Icons.BAR_CHART, size=14, color=Colors.FG_DIM),
                                    ft.Text(label, size=11, color=Colors.FG, expand=True),
                                ]),
                                padding=ft.padding.Padding(left=28, top=3, right=6, bottom=3),
                                border_radius=6, bgcolor=Colors.TREE_BG,
                                on_click=make_click(), ink=True,
                            ))
            page.update()

        rebuild()

        dlg = ft.AlertDialog(
            title=ft.Text("Отчёты", size=18, weight=ft.FontWeight.BOLD, color=Colors.ACCENT),
            content=ft.Container(content=content_col, width=360, height=450),
            actions=[ft.TextButton("Закрыть", on_click=lambda _: page.pop_dialog())],
        )
        page.show_dialog(dlg)

    def show_report(title, headers, rows):
        cols = [ft.DataColumn(ft.Text(h, size=11, weight=ft.FontWeight.BOLD, color=Colors.FG)) for h in headers]
        tbl_rows = []
        for row in rows:
            cells = [ft.DataCell(ft.Text(str(v), size=10, color=Colors.FG)) for v in row]
            tbl_rows.append(ft.DataRow(cells=cells))

        tbl = ft.DataTable(columns=cols, rows=tbl_rows,
                           heading_row_color=Colors.BG_LIGHT,
                           data_row_color={"": Colors.TREE_BG},
                           divider_thickness=0.5,
                           border=ft.Border(
            left=ft.BorderSide(0.5, Colors.BG_LIGHT),
            top=ft.BorderSide(0.5, Colors.BG_LIGHT),
            right=ft.BorderSide(0.5, Colors.BG_LIGHT),
            bottom=ft.BorderSide(0.5, Colors.BG_LIGHT),
        ),
                           column_spacing=8)

        def export_excel(_):
            try:
                import openpyxl
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = title[:31]
                ws.append(headers)
                for row in rows:
                    ws.append([str(v) for v in row])
                path = os.path.join(APP_DIR, f"{title}.xlsx")
                wb.save(path)
                show_snack(f"Excel: {path}")
            except Exception as ex:
                show_snack(f"Ошибка: {ex}", Colors.ERROR)

        def export_pdf(_):
            try:
                from fpdf import FPDF
                arial = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts", "arial.ttf")
                arial_b = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts", "arialbd.ttf")
                pdf = FPDF(orientation="L", unit="mm", format="A4")
                pdf.add_font("Arial", "", arial, uni=True)
                pdf.add_font("Arial", "B", arial_b, uni=True)
                pdf.add_page()
                pdf.set_font("Arial", size=14)
                pdf.cell(0, 10, title, ln=True, align="C")
                pdf.ln(4)
                num_cols = len(headers)
                page_w = pdf.w - 2 * pdf.l_margin
                col_w = page_w / num_cols
                pdf.set_font("Arial", size=10, style="B")
                for h in headers:
                    pdf.cell(col_w, 8, str(h), border=1, align="C")
                pdf.ln()
                pdf.set_font("Arial", size=9)
                for row in rows:
                    for val in row:
                        pdf.cell(col_w, 7, str(val), border=1, align="C")
                    pdf.ln()
                path = os.path.join(APP_DIR, f"{title}.pdf")
                pdf.output(path)
                show_snack(f"PDF: {path}")
            except Exception as ex:
                show_snack(f"Ошибка: {ex}", Colors.ERROR)

        report_dlg = ft.AlertDialog(
            title=ft.Text(title, size=16, weight=ft.FontWeight.BOLD, color=Colors.ACCENT),
            content=ft.Container(
                content=ft.Column([
                    ft.Container(content=tbl, padding=4),
                    ft.Row([
                        ft.Button("Excel", icon=ft.Icons.TABLE_CHART,
                                          on_click=export_excel, bgcolor=Colors.BG_LIGHT,
                                          color=Colors.FG),
                        ft.Button("PDF", icon=ft.Icons.PICTURE_AS_PDF,
                                          on_click=export_pdf, bgcolor=Colors.BG_LIGHT,
                                          color=Colors.FG),
                    ], alignment=ft.MainAxisAlignment.CENTER),
                ], scroll=ft.ScrollMode.AUTO),
                width=380,
                height=500,
            ),
            actions=[ft.TextButton("Закрыть", on_click=lambda _: page.pop_dialog())],
        )
        page.show_dialog(report_dlg)

    def report_revenue_total():
        r = db.query("SELECT SUM(price) AS total FROM tickets")
        val = r[0]["total"] or 0
        show_report("Общая выручка", ["Метрика", "Значение"],
                    [["Общая выручка", f"{val:,.2f} ₽"]])

    def report_revenue_by_type():
        r = db.query("SELECT ticket_type, SUM(price) AS total FROM tickets GROUP BY ticket_type")
        show_report("Выручка по типу билета", ["Тип", "Выручка (₽)"],
                    [[x[0], f"{x[1]:,.2f}"] for x in [list(row) for row in r]])

    def report_avg_price():
        r = db.query("SELECT AVG(price) AS avg_p FROM tickets")
        val = r[0]["avg_p"] or 0
        show_report("Средний чек", ["Метрика", "Значение"], [["Средний чек", f"{val:,.2f} ₽"]])

    def report_by_ticket_type():
        r = db.query("SELECT ticket_type, COUNT(*) AS cnt FROM tickets GROUP BY ticket_type ORDER BY cnt DESC")
        show_report("Посетители по типу билета", ["Тип", "Кол-во"], [list(row) for row in r])

    def report_by_age():
        r = db.query("""
            SELECT CASE
                WHEN age < 7 THEN '0-6'
                WHEN age < 14 THEN '7-13'
                WHEN age < 18 THEN '14-17'
                WHEN age < 30 THEN '18-29'
                WHEN age < 50 THEN '30-49'
                ELSE '50+'
            END AS age_group, COUNT(*) AS cnt
            FROM tickets GROUP BY age_group ORDER BY age_group
        """)
        show_report("Посетители по возрасту", ["Группа", "Кол-во"], [list(row) for row in r])

    def report_popularity():
        r = db.query("SELECT attraction, COUNT(*) AS cnt FROM tickets GROUP BY attraction ORDER BY cnt DESC")
        show_report("Популярность аттракционов", ["Аттракцион", "Посещений"], [list(row) for row in r])

    def report_attraction_revenue():
        r = db.query("SELECT attraction, SUM(price) AS total FROM tickets GROUP BY attraction ORDER BY total DESC")
        show_report("Доход по аттракционам", ["Аттракцион", "Доход (₽)"],
                    [[x[0], f"{x[1]:,.2f}"] for x in [list(row) for row in r]])

    def report_sales_by_date():
        r = db.query("SELECT purchase_date, COUNT(*) AS cnt, SUM(price) AS total FROM tickets GROUP BY purchase_date ORDER BY purchase_date DESC")
        show_report("Продажи по датам", ["Дата", "Билетов", "Сумма (₽)"],
                    [[x[0], x[1], f"{x[2]:,.2f}"] for x in [list(row) for row in r]])

    def report_expensive():
        r = db.query("SELECT * FROM tickets ORDER BY price DESC LIMIT 5")
        show_report("Топ-5 дорогих билетов",
                    ["ID", "Аттракцион", "Посетитель", "Возраст", "Цена", "Тип", "Дата"],
                    [list(row) for row in r])

    def report_top_visitors():
        r = db.query("SELECT visitor_name, SUM(price) AS total, COUNT(*) AS visits FROM tickets GROUP BY visitor_name ORDER BY total DESC LIMIT 5")
        show_report("Топ посетителей", ["Посетитель", "Сумма (₽)", "Визитов"],
                    [[x[0], f"{x[1]:,.2f}", x[2]] for x in [list(row) for row in r]])

    def report_price_range():
        r = db.query("SELECT MIN(price), MAX(price), ROUND(AVG(price),1) FROM tickets")
        x = r[0]
        show_report("Ценовой диапазон", ["Метрика", "Значение"],
                    [["Мин. цена", f"{x[0]:,.2f} ₽"], ["Макс. цена", f"{x[1]:,.2f} ₽"],
                      ["Средняя цена", f"{x[2]:,.2f} ₽"]])

    report_tree = {
        "1. Финансовые отчёты": [
            ("Общая выручка", report_revenue_total),
            ("Выручка по типу билета", report_revenue_by_type),
            ("Средний чек", report_avg_price),
        ],
        "2. Посещаемость": [
            ("Посетители по типу билета", report_by_ticket_type),
            ("Посетители по возрасту", report_by_age),
        ],
        "3. Аттракционы": [
            ("Популярность аттракционов", report_popularity),
            ("Доход по аттракционам", report_attraction_revenue),
        ],
        "4. Билеты": [
            ("Продажи по датам", report_sales_by_date),
            ("Дорогие билеты", report_expensive),
        ],
        "5. Аналитика": [
            ("Топ посетителей", report_top_visitors),
            ("Ценовой диапазон", report_price_range),
        ],
    }

    # ─── Об авторе ───
    def show_about(_=None):
        dlg = ft.AlertDialog(
            title=ft.Text("Freddy Fazbear DB", size=18, weight=ft.FontWeight.BOLD, color=Colors.ACCENT),
            content=ft.Column([
                ft.Text("Автор: Кабушко Максим Николаевич", size=14, weight=ft.FontWeight.BOLD),
                ft.Text("Группа: ИС-943  |  Курс: 2", size=12),
                ft.Divider(color=Colors.BG_LIGHT),
                ft.Text("Библиотеки:", size=12, weight=ft.FontWeight.BOLD),
                ft.Text("  • flet — UI", size=12, color=Colors.FG_DIM),
                ft.Text("  • sqlite3 — реляционная БД", size=12, color=Colors.FG_DIM),
                ft.Text("  • openpyxl — экспорт Excel", size=12, color=Colors.FG_DIM),
                ft.Text("  • fpdf2 — экспорт PDF", size=12, color=Colors.FG_DIM),
            ], spacing=4),
            actions=[ft.TextButton("Закрыть", on_click=lambda _: page.pop_dialog())],
        )
        page.show_dialog(dlg)

    # ─── Нижняя панель ───
    bottom_bar = ft.Container(
        content=ft.Row(
            [
                ft.Button("Добавить", icon=ft.Icons.ADD, on_click=open_add_dialog,
                                  bgcolor=Colors.ACCENT, color=Colors.WHITE,
                                  style=ft.ButtonStyle(shape=ft.RoundedRectangleBorder(radius=8), padding=12)),
                ft.Button("Удалить", icon=ft.Icons.DELETE, on_click=delete_selected,
                                  bgcolor=Colors.BG_LIGHT, color=Colors.FG,
                                  style=ft.ButtonStyle(shape=ft.RoundedRectangleBorder(radius=8), padding=12)),
                ft.Button("Отчёты", icon=ft.Icons.BAR_CHART, on_click=open_reports,
                                  bgcolor=Colors.BG_LIGHT, color=Colors.FG,
                                  style=ft.ButtonStyle(shape=ft.RoundedRectangleBorder(radius=8), padding=12)),
                ft.IconButton(ft.Icons.INFO_OUTLINE, icon_color=Colors.FG_DIM,
                              on_click=show_about, icon_size=22),
            ],
            alignment=ft.MainAxisAlignment.SPACE_EVENLY,
        ),
        bgcolor=Colors.BG_MID,
        padding=ft.padding.Padding(left=8, top=8, right=8, bottom=8),
        border_radius=ft.BorderRadius.only(top_left=12, top_right=12),
    )

    # ─── Сборка страницы ───
    rebuild_table()

    page.add(
        ft.Column(
            [
                header,
                search_bar,
                col_filter_bar,
                table_scroll,
                ft.Container(
                    content=ft.Row([status_text, count_text], alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
                    padding=ft.padding.Padding(left=12, top=4, right=12, bottom=4),
                ),
                bottom_bar,
            ],
            expand=True,
            spacing=0,
        )
    )


if __name__ == "__main__":
    ft.run(main)
